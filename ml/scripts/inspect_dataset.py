#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
RiceLeafDiseaseBD Dataset Inspector - KrishiSathi AI (Phase 1)

Run after download_dataset.py. Checks, in order:
  1. Folder counts vs. the official metadata sheet
  2. Every image decodes; size / mode
  3. Exact duplicates (MD5) and near-duplicates (64-bit dHash)
  4. Same source photo filed under two classes (label conflicts)
  5. Capture-time bursts from original filenames (leakage risk for splits)
  6. Shortcut risks: filename pattern / original resolution / brightness per class
  7. YOLO label files vs. the README class-ID mapping

Outputs (in ml/reports/):
  dataset_inspection.txt          - human-readable report (also printed)
  manifest.csv                    - one row per image; input for the split step
  near_duplicates_cross_class.csv - pairs to review by eye

Usage:
  python ml/scripts/inspect_dataset.py ml/data/raw/riceleafdiseasebd
"""

import argparse
import csv
import hashlib
import io
import re
import sys
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from PIL import Image

# README class-ID mapping for the YOLO label files.
README_CLASS_IDS = {'Brown spot': 0, 'Blast': 1, 'Healthy': 2, 'Leaf smut': 3,
                    'Rice Tungro': 4, 'Sheath blight': 5}
BD_TZ = timezone(timedelta(hours=6))
# Photos closer than this in time = same burst. Measured on v3: consecutive photos
# <2s apart are near-copies (median dHash 10), from 5s on they look like random
# pairs (median ~30). 60s chained whole sessions into one 900+ image group.
BURST_GAP_S = 10
NEAR_DUP_BITS = 4      # dHash Hamming distance treated as near-duplicate
POPCOUNT = np.array([bin(i).count('1') for i in range(256)], dtype=np.uint8)


class Report:
    def __init__(self):
        self.lines = []

    def __call__(self, *parts):
        line = ' '.join(str(p) for p in parts)
        print(line)
        self.lines.append(line)

    def section(self, title):
        self('')
        self(f'## {title}')


# ---------------------------------------------------------------- metadata

def capture_time_utc(original_name):
    """Best-effort capture time from the camera's original filename.

    Assumptions (documented in the report): 13-digit names are epoch
    milliseconds; Pixel 'PXL_' names are UTC; other date-stamped names are
    Bangladesh local time (UTC+6).
    """
    m = re.match(r'^(\d{13})\.', original_name)
    if m:
        return datetime.fromtimestamp(int(m[1]) / 1000, tz=timezone.utc), 'epoch_ms'
    m = re.match(r'^PXL_(\d{8})_(\d{6})', original_name)
    if m:
        return datetime.strptime(m[1] + m[2], '%Y%m%d%H%M%S').replace(tzinfo=timezone.utc), 'PXL_'
    m = re.match(r'^IMG(\d{8})(\d{6})', original_name)
    if m:
        return datetime.strptime(m[1] + m[2], '%Y%m%d%H%M%S').replace(tzinfo=BD_TZ), 'IMGyyyymmdd'
    m = re.match(r'^(\d{8})_(\d{6})', original_name)
    if m:
        return datetime.strptime(m[1] + m[2], '%Y%m%d%H%M%S').replace(tzinfo=BD_TZ), 'yyyymmdd_hhmmss'
    return None, 'unknown'


def load_metadata(path):
    ws = load_workbook(path, read_only=True)['Image_Metadata']
    rows = ws.iter_rows(values_only=True)
    header = list(next(rows))
    meta = {}
    for row in rows:
        r = dict(zip(header, row))
        if r.get('New_Filename'):
            meta[r['New_Filename']] = r
    return meta


# ---------------------------------------------------------------- images

def dhash64(gray):
    small = np.asarray(gray.resize((9, 8), Image.BILINEAR), dtype=np.int16)
    bits = (small[:, 1:] > small[:, :-1]).flatten()
    return int(''.join('1' if b else '0' for b in bits), 2)


def analyze_image(path):
    data = path.read_bytes()
    out = {'md5': hashlib.md5(data).hexdigest(), 'error': ''}
    try:
        img = Image.open(io.BytesIO(data))
        img.load()  # full decode: catches truncated/corrupt files
        out.update(width=img.width, height=img.height, mode=img.mode)
        gray = img.convert('L')
        out['dhash'] = dhash64(gray)
        g = np.asarray(gray.resize((256, 256), Image.BILINEAR), dtype=np.float32)
        lap = (g[1:-1, :-2] + g[1:-1, 2:] + g[:-2, 1:-1] + g[2:, 1:-1] - 4 * g[1:-1, 1:-1])
        out['brightness'] = float(g.mean())
        out['blur_var'] = float(lap.var())  # low = blurry
    except Exception as e:  # noqa: BLE001 - report every failure, keep going
        out['error'] = f'{type(e).__name__}: {e}'
    return out


def near_duplicate_pairs(hashes, max_bits):
    h = np.array(hashes, dtype=np.uint64)
    pairs = []
    for start in range(0, len(h), 256):
        block = h[start:start + 256, None] ^ h[None, :]
        dist = POPCOUNT[block.view(np.uint8)].reshape(block.shape + (8,)).sum(-1)
        ii, jj = np.nonzero(dist <= max_bits)
        for i, j in zip(ii + start, jj):
            if i < j:
                pairs.append((int(i), int(j), int(dist[i - start, j])))
    return pairs


class UnionFind:
    def __init__(self, n):
        self.p = list(range(n))

    def find(self, x):
        while self.p[x] != x:
            self.p[x] = self.p[self.p[x]]
            x = self.p[x]
        return x

    def union(self, a, b):
        self.p[self.find(a)] = self.find(b)


# ---------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('root', help='folder created by download_dataset.py')
    ap.add_argument('--reports', default='ml/reports')
    args = ap.parse_args()

    root = Path(args.root)
    img_root = root / 'Original images'
    meta_path = root / 'Dataset metadata.xlsx'
    if not img_root.is_dir() or not meta_path.exists():
        sys.exit(f'[ERROR] Expected "{img_root}" and "{meta_path}". Run download_dataset.py first.')
    reports = Path(args.reports)
    reports.mkdir(parents=True, exist_ok=True)

    R = Report()
    R('=' * 70)
    R('RICELEAFDISEASEBD v3 INSPECTION - KrishiSathi AI (Phase 1)')
    R('=' * 70)
    R(f'Root: {root.resolve()}')
    R(f'Run : {datetime.now().isoformat(timespec="seconds")}')

    # 1) counts vs metadata -------------------------------------------------
    files = sorted(p for p in img_root.glob('*/*') if p.is_file() and not p.name.endswith('.part'))
    meta = load_metadata(meta_path)
    R.section('1) CLASS COUNTS (folders vs metadata sheet)')
    folder_counts = Counter(p.parent.name for p in files)
    meta_counts = Counter()
    for r in meta.values():
        meta_counts[r['Class']] += 1
    meta_lower = {k.lower(): v for k, v in meta_counts.items()}
    for cls, n in sorted(folder_counts.items(), key=lambda kv: -kv[1]):
        expected = meta_lower.get(cls.lower(), 0)
        flag = '' if n == expected else f'   <-- metadata says {expected}'
        R(f'  {cls:15s} {n:5d}{flag}')
    total = sum(folder_counts.values())
    R(f'  {"TOTAL":15s} {total:5d}   (metadata: {sum(meta_counts.values())})')
    exts = Counter(p.suffix.lower() for p in files)
    R(f'  extensions: {dict(exts)}')
    missing_meta = [p.name for p in files if p.name not in meta]
    R(f'  images without a metadata row: {len(missing_meta)} {missing_meta[:5]}')
    biggest, smallest = folder_counts.most_common()[0], folder_counts.most_common()[-1]
    R(f'  imbalance max/min: {biggest[0]} {biggest[1]} / {smallest[0]} {smallest[1]}'
      f' = {biggest[1] / smallest[1]:.2f}x')

    # 2) decode every image -------------------------------------------------
    R.section('2) IMAGE DECODE / SIZE')
    print(f'  decoding {len(files)} images ...', flush=True)
    with ThreadPoolExecutor(max_workers=8) as pool:
        stats = list(pool.map(analyze_image, files))
    bad = [(p, s['error']) for p, s in zip(files, stats) if s['error']]
    R(f'  failed to decode: {len(bad)}')
    for p, err in bad[:10]:
        R(f'    - {p.relative_to(root)}: {err}')
    ok = [(p, s) for p, s in zip(files, stats) if not s['error']]
    R(f'  sizes: {dict(Counter((s["width"], s["height"]) for _, s in ok))}')
    R(f'  modes: {dict(Counter(s["mode"] for _, s in ok))}')

    # 3) duplicates ---------------------------------------------------------
    R.section('3) DUPLICATES')
    by_md5 = defaultdict(list)
    for p, s in ok:
        by_md5[s['md5']].append(p)
    exact = [g for g in by_md5.values() if len(g) > 1]
    R(f'  exact (MD5) duplicate groups: {len(exact)}')
    for g in exact[:10]:
        R('    - ' + ' == '.join(str(p.relative_to(img_root)) for p in g))

    pairs = near_duplicate_pairs([s['dhash'] for _, s in ok], max_bits=8)
    for bits in (0, 2, 4, 6, 8):
        n_pairs = sum(1 for _, _, d in pairs if d <= bits)
        n_cross = sum(1 for i, j, d in pairs if d <= bits and ok[i][0].parent != ok[j][0].parent)
        R(f'  dHash distance <= {bits}: {n_pairs:6d} pairs, {n_cross:5d} across different classes')
    cross = [(i, j, d) for i, j, d in pairs if d <= NEAR_DUP_BITS and ok[i][0].parent != ok[j][0].parent]
    with open(reports / 'near_duplicates_cross_class.csv', 'w', newline='') as f:
        w = csv.writer(f)
        w.writerow(['image_a', 'image_b', 'dhash_distance'])
        for i, j, d in cross:
            w.writerow([ok[i][0].relative_to(img_root), ok[j][0].relative_to(img_root), d])
    R(f'  cross-class near-duplicates (<= {NEAR_DUP_BITS} bits) written to near_duplicates_cross_class.csv'
      ' -> open each pair and decide which label is right, or drop both')

    # 4) same source photo in two classes -----------------------------------
    R.section('4) SAME ORIGINAL PHOTO FILED UNDER MORE THAN ONE CLASS')
    by_orig = defaultdict(list)
    for name, r in meta.items():
        by_orig[r['Original_Filename']].append(name)
    conflicts = {k: v for k, v in by_orig.items() if len(v) > 1}
    R(f'  {len(conflicts)} original filenames appear more than once:')
    for k, v in conflicts.items():
        R(f'    - {k}: {v}')
    R('  (filenames can collide across phones; check visually before deleting)')

    # 5) capture bursts -----------------------------------------------------
    R.section(f'5) CAPTURE-TIME BURSTS (gap <= {BURST_GAP_S}s = same burst)')
    rows = []
    for idx, (p, s) in enumerate(ok):
        m = meta.get(p.name, {})
        t, pattern = capture_time_utc(str(m.get('Original_Filename', '')))
        rows.append({'idx': idx, 'path': p, 'cls': p.parent.name, 'stats': s, 'meta': m,
                     'time': t, 'pattern': pattern})
    # Bursts are formed per class: two people often shot different diseased
    # leaves side by side, and chaining across classes merged 400+ images.
    timed = sorted((r for r in rows if r['time']), key=lambda r: (r['cls'], r['time']))
    R(f'  capture time parsed: {len(timed)}/{len(rows)}')
    R('  ASSUMPTION: epoch-ms and PXL_ names are UTC; other dated names are UTC+6 local time')
    burst = {}
    b = 0
    for prev, cur in zip([None] + timed[:-1], timed):
        if (prev is None or prev['cls'] != cur['cls']
                or (cur['time'] - prev['time']).total_seconds() > BURST_GAP_S):
            b += 1
        burst[cur['idx']] = f'b{b:05d}'
    for r in rows:
        r['burst'] = burst.get(r['idx'], f'solo{r["idx"]:05d}')
    burst_sizes = Counter(r['burst'] for r in rows)
    R(f'  distinct bursts: {len(burst_sizes)} for {len(rows)} images'
      f' (median {int(np.median(list(burst_sizes.values())))} img/burst,'
      f' largest {max(burst_sizes.values())})')
    per_cls = defaultdict(set)
    for r in rows:
        per_cls[r['cls']].add(r['burst'])
    for cls in sorted(per_cls):
        R(f'    {cls:15s} {len(per_cls[cls]):4d} bursts')
    mixed = sum(1 for bid in burst_sizes if len({r['cls'] for r in rows if r['burst'] == bid}) > 1)
    R(f'  bursts containing more than one class: {mixed}')
    days = Counter(r['time'].astimezone(BD_TZ).date().isoformat() for r in timed)
    R(f'  capture days (BD time): {len(days)}  {min(days)} .. {max(days)}')

    # Final leakage group = burst, merged with near-duplicates and shared originals.
    uf = UnionFind(len(rows))
    first_in_burst = {}
    for r in rows:
        first_in_burst.setdefault(r['burst'], r['idx'])
        uf.union(r['idx'], first_in_burst[r['burst']])
    for i, j, d in pairs:
        if d <= NEAR_DUP_BITS:
            uf.union(i, j)
    name_to_idx = {r['path'].name: r['idx'] for r in rows}
    for names in conflicts.values():
        ids = [name_to_idx[n] for n in names if n in name_to_idx]
        for a in ids[1:]:
            uf.union(ids[0], a)
    roots = {}
    for r in rows:
        r['group'] = roots.setdefault(uf.find(r['idx']), f'g{len(roots):05d}')
    group_sizes = Counter(r['group'] for r in rows)
    R(f'  leakage groups (burst + near-dup + shared original): {len(group_sizes)},'
      f' largest {max(group_sizes.values())} images')
    R('  -> train/val/test MUST be split by group, never image-by-image')

    # 6) shortcut risks -----------------------------------------------------
    R.section('6) SHORTCUT RISKS (things a model could learn instead of disease)')
    R('  original filename pattern per class (proxy for phone / transfer path):')
    pats = sorted({r['pattern'] for r in rows})
    R('    ' + ' ' * 15 + ''.join(f'{p:>17s}' for p in pats))
    for cls in sorted(per_cls):
        c = Counter(r['pattern'] for r in rows if r['cls'] == cls)
        n = sum(c.values())
        R(f'    {cls:15s}' + ''.join(f'{c[p]:>10d} ({100 * c[p] / n:3.0f}%)' for p in pats))
    R('  original resolution per class (top 3):')
    for cls in sorted(per_cls):
        c = Counter(f'{int(r["meta"].get("Original_Width_px") or 0)}x{int(r["meta"].get("Original_Height_px") or 0)}'
                    for r in rows if r['cls'] == cls)
        R(f'    {cls:15s} ' + ', '.join(f'{k}: {v}' for k, v in c.most_common(3)))
    R('  brightness / sharpness per class (mean, [p5..p95]); blur_var low = blurry:')
    for cls in sorted(per_cls):
        br = np.array([r['stats']['brightness'] for r in rows if r['cls'] == cls])
        bl = np.array([r['stats']['blur_var'] for r in rows if r['cls'] == cls])
        R(f'    {cls:15s} brightness {br.mean():6.1f} [{np.percentile(br, 5):5.1f}..{np.percentile(br, 95):5.1f}]'
          f'   blur_var {np.median(bl):7.1f} [{np.percentile(bl, 5):6.1f}..{np.percentile(bl, 95):7.1f}]')

    # 7) YOLO labels ----------------------------------------------------------
    R.section('7) YOLO LABEL FILES (README mapping: ' +
              ', '.join(f'{v}={k}' for k, v in sorted(README_CLASS_IDS.items(), key=lambda kv: kv[1])) + ')')
    label_root = root / 'labels'
    if not label_root.is_dir():
        R('  labels/ folder not found - skipped')
    else:
        for cls_dir in sorted(p for p in label_root.iterdir() if p.is_dir()):
            ids, n_files, n_boxes, empty, bad_coord = Counter(), 0, 0, 0, 0
            for f in cls_dir.glob('*.txt'):
                n_files += 1
                lines = [ln.split() for ln in f.read_text().splitlines() if ln.strip()]
                empty += not lines
                for parts in lines:
                    n_boxes += 1
                    ids[parts[0]] += 1
                    vals = [float(x) for x in parts[1:5]]
                    bad_coord += len(parts) != 5 or not all(0 <= v <= 1 for v in vals)
            expected = README_CLASS_IDS.get(cls_dir.name)
            seen = dict(ids)
            flag = '' if set(seen) == {str(expected)} else f'   <-- README says {expected}'
            R(f'  {cls_dir.name:15s} files {n_files:5d}  boxes {n_boxes:6d}  empty {empty:4d}'
              f'  bad_coords {bad_coord}  class_ids {seen}{flag}')
        R('  Classification does not use these. If they are used later (detection / lesion crops),')
        R('  remap IDs from the FOLDER name, not from the number inside the file.')

    # manifest ----------------------------------------------------------------
    with open(reports / 'manifest.csv', 'w', newline='') as f:
        w = csv.writer(f)
        w.writerow(['path', 'class', 'original_filename', 'filename_pattern', 'capture_time_utc',
                    'burst_id', 'group_id', 'md5', 'dhash', 'width', 'height',
                    'brightness', 'blur_var'])
        for r in rows:
            s = r['stats']
            w.writerow([r['path'].relative_to(root), r['cls'], r['meta'].get('Original_Filename', ''),
                        r['pattern'], r['time'].astimezone(timezone.utc).isoformat() if r['time'] else '',
                        r['burst'], r['group'],
                        s['md5'], f'{s["dhash"]:016x}', s['width'], s['height'],
                        f'{s["brightness"]:.1f}', f'{s["blur_var"]:.1f}'])

    R('')
    R('=' * 70)
    R('Written: ml/reports/dataset_inspection.txt, manifest.csv, near_duplicates_cross_class.csv')
    R('Next  : review section 3/4 pairs by eye, then share this report to plan the split.')
    R('=' * 70)
    (reports / 'dataset_inspection.txt').write_text('\n'.join(R.lines) + '\n', encoding='utf-8')


if __name__ == '__main__':
    main()
